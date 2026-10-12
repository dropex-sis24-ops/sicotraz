from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
IMAGES = ROOT / "evidencia_capturas"
OUTPUT = ROOT / "Evidencia_Avance_SICOTRAZ_UT5.docx"


def set_run_font(run, size=10, bold=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_paragraph(paragraph, size=10, bold=False, alignment=None, after=6):
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.05
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold)


def add_label_value(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    label_run = paragraph.add_run(label + ": ")
    set_run_font(label_run, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run)


def add_evidence_page(doc, number, title, description, filename):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    heading = doc.add_paragraph()
    run = heading.add_run(f"{number}. {title}")
    set_run_font(run, size=12, bold=True)
    heading.paragraph_format.space_after = Pt(6)

    body = doc.add_paragraph(description)
    style_paragraph(body, size=10, after=8)

    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_run = picture_paragraph.add_run()
    picture_run.add_picture(str(IMAGES / filename), width=Inches(7.0))
    picture_paragraph.paragraph_format.space_after = Pt(4)

    caption = doc.add_paragraph(f"Captura {number}. {title} del sistema SICOTRAZ.")
    style_paragraph(caption, size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(32)
title.paragraph_format.space_after = Pt(10)
title_run = title.add_run("EVIDENCIA DEL SISTEMA DESARROLLADO")
set_run_font(title_run, size=14, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(28)
subtitle_run = subtitle.add_run("Avance parcial de la UT5")
set_run_font(subtitle_run, size=11, bold=True)

add_label_value(doc, "Nombre del estudiante", "Pedro Jesus Chucamani Choque")
add_label_value(doc, "Nombre del proyecto", "SICOTRAZ")
add_label_value(doc, "Fecha", "28 de septiembre de 2026")
add_label_value(doc, "Estado del desarrollo", "Avance intermedio aproximado del 50 %")

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(22)
intro.paragraph_format.space_after = Pt(8)
intro_run = intro.add_run("Descripción del avance")
set_run_font(intro_run, size=11, bold=True)

paragraph = doc.add_paragraph(
    "El sistema se encuentra en una etapa intermedia de desarrollo. Las capturas siguientes "
    "evidencian las funciones básicas que ya están implementadas y operativas: acceso al sistema, "
    "menú principal, administración inicial de usuarios, consulta de catálogos, preparación de "
    "plantillas y visualización del tablero de resumen. Los demás procesos continúan en desarrollo, "
    "integración y revisión."
)
style_paragraph(paragraph, size=10, after=8)

scope = doc.add_paragraph(
    "La evidencia corresponde a una ejecución local del proyecto y utiliza información de prueba."
)
style_paragraph(scope, size=10, after=0)

add_evidence_page(
    doc,
    1,
    "Inicio de sesión",
    "Pantalla de acceso mediante número de ítem o contrato y contraseña. El botón de ingreso se habilita cuando se completan ambos campos.",
    "01_inicio_sesion.png",
)
add_evidence_page(
    doc,
    2,
    "Pantalla principal",
    "Menú principal disponible después de una autenticación correcta. Desde esta pantalla se accede a las funciones habilitadas según el rol del usuario.",
    "02_pantalla_principal.png",
)
add_evidence_page(
    doc,
    3,
    "Gestión de usuarios",
    "Listado de usuarios registrados con nombre, rol, número de ítem y estado. También se observa el campo de búsqueda y la opción para registrar un nuevo usuario.",
    "03_gestion_usuarios.png",
)
add_evidence_page(
    doc,
    4,
    "Formulario de nuevo usuario",
    "Formulario implementado para registrar usuarios mediante nombre completo, número de ítem, carnet de identidad y rol. En esta evidencia no se guardó ningún registro adicional.",
    "04_formulario_nuevo_usuario.png",
)
add_evidence_page(
    doc,
    5,
    "Catálogo de prendas",
    "Listado de tipos de prendas disponibles en el sistema, con indicación de su estado y opciones básicas de administración.",
    "05_catalogo_prendas.png",
)
add_evidence_page(
    doc,
    6,
    "Catálogo de áreas y alias",
    "Sección destinada a la consulta y mantenimiento inicial de las áreas de trabajo y sus denominaciones alternativas.",
    "06_catalogo_areas.png",
)
add_evidence_page(
    doc,
    7,
    "Catálogo de plantillas",
    "Configuración de las plantillas habilitadas para Salas y Quirófano, utilizadas posteriormente en los procesos de registro.",
    "07_catalogo_plantillas.png",
)
add_evidence_page(
    doc,
    8,
    "Generación de plantillas en blanco",
    "Pantalla para seleccionar el tipo de plantilla, indicar opcionalmente un área preimpresa y generar el archivo para impresión.",
    "08_impresion_plantillas.png",
)
add_evidence_page(
    doc,
    9,
    "Dashboard de seguimiento",
    "Tablero de resumen que presenta indicadores básicos de alertas, prendas procesadas, bajas y circulación de ropa. Los valores corresponden al entorno local de prueba.",
    "09_dashboard.png",
)

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        if run.font.name is None:
            set_run_font(run)

doc.core_properties.title = "Evidencia del sistema desarrollado SICOTRAZ"
doc.core_properties.subject = "Avance parcial UT5"
doc.core_properties.author = "Pedro Jesus Chucamani Choque"
doc.save(OUTPUT)
print(OUTPUT)
