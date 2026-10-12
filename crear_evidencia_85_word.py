from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "evidencia_capturas"
OUTPUT = ROOT / "Evidencia_Avance_85_SICOTRAZ.docx"


def set_font(run, size=10, bold=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def format_paragraph(paragraph, size=10, bold=False, alignment=None, after=6, line=1.08):
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8 if level == 1 else 5)
    r = p.add_run(text)
    set_font(r, size=13 if level == 1 else 11, bold=True)
    return p


def add_body(doc, text, after=7):
    p = doc.add_paragraph(text)
    format_paragraph(p, size=10, after=after)
    return p


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    set_font(run, size=8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor(0, 0, 0)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.text = "SICOTRAZ - Evidencia de avance del proyecto"
format_paragraph(footer_p, size=8, after=0)
add_page_number(footer.add_paragraph())

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(42)
title.paragraph_format.space_after = Pt(10)
set_font(title.add_run("EVIDENCIA DE AVANCE DEL PROYECTO"), size=16, bold=True)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(24)
set_font(subtitle.add_run("Registro de objetivos actividades y resultados"), size=11, bold=True)

meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.style = "Table Grid"
metadata = [
    ("Nombre del proyecto", "SICOTRAZ - Sistema para el control, digitalización y trazabilidad de ropa hospitalaria"),
    ("Nombre del estudiante", "Pedro Jesus Chucamani Choque"),
    ("Fecha", "11 de octubre de 2026"),
    ("Avance declarado", "85 %"),
]
for row, (label, value) in zip(meta.rows, metadata):
    row.cells[0].width = Inches(1.6)
    row.cells[1].width = Inches(5.3)
    set_cell_shading(row.cells[0], "EDEDED")
    for cell in row.cells:
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[0].text = label
    row.cells[1].text = value
    format_paragraph(row.cells[0].paragraphs[0], size=10, bold=True, after=0)
    format_paragraph(row.cells[1].paragraphs[0], size=10, after=0)

doc.add_paragraph()
add_heading(doc, "Resumen del avance")
add_body(doc, "El proyecto alcanzó un avance estimado del 85 %. Se encuentran implementados los módulos principales de autenticación, usuarios, catálogos, stock, movimientos, alertas, OCR local, bajas, dashboard, reportes y sincronización. La auditoría técnica permitió corregir problemas críticos de seguridad, integridad de datos, concurrencia y funcionamiento sin conexión.")
add_body(doc, "El porcentaje restante corresponde a validación de campo con usuarios institucionales, ajustes derivados de esa validación, preparación del entorno de producción y decisiones de producto documentadas en la auditoría.")

doc.add_page_break()
add_heading(doc, "1 Objetivos registrados")
add_heading(doc, "Objetivo general", level=2)
add_body(doc, "Desarrollar un sistema para el control, digitalización y trazabilidad de movimientos de ropa hospitalaria en el Hospital de Tercer Nivel, mediante tecnología OCR, con el propósito de optimizar la gestión de información y fortalecer el control de los recursos textiles institucionales.")
add_heading(doc, "Objetivos específicos", level=2)
specific = [
    "Analizar el proceso actual de gestión de ropa hospitalaria e identificar oportunidades de mejora y necesidades de control.",
    "Diseñar la arquitectura, la base de datos y las interfaces del sistema para una gestión eficiente de la información.",
    "Implementar un módulo de digitalización de formularios mediante OCR para automatizar el registro de recepción y entrega de prendas.",
    "Desarrollar funciones de control, consulta y trazabilidad para registrar y monitorear los movimientos de ropa hospitalaria.",
    "Generar reportes e indicadores que apoyen el control y la toma de decisiones.",
    "Validar el funcionamiento mediante pruebas operativas de registro, digitalización, consulta y seguimiento.",
]
for item in specific:
    p = doc.add_paragraph(style=None)
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("•  " + item)
    set_font(r)

doc.add_page_break()
add_heading(doc, "2 Relación entre objetivos actividades y avances")
matrix_data = [
    ("OE1", "Levantamiento y análisis del flujo de lavandería y ropería.", "Requerimientos, reglas de negocio y backlog documentados.", "Completado"),
    ("OE2", "Diseño de arquitectura Laravel y Flutter, base de datos y pantallas por rol.", "API, aplicación móvil, migraciones y catálogos implementados.", "Avanzado"),
    ("OE3", "Implementación y mejora del OCR local.", "Lectura, normalización, revisión editable, detección de ambigüedades y 8 pruebas OCR.", "Avanzado"),
    ("OE4", "Desarrollo de stock, lotes, etapas, alertas, bajas y sincronización.", "Trazabilidad funcional con controles de concurrencia, idempotencia y permisos por área.", "Avanzado"),
    ("OE5", "Construcción de dashboard, reportes y plantillas PDF.", "Indicadores y salidas disponibles; quedan mejoras documentadas.", "Parcial alto"),
    ("OE6", "Auditoría técnica, corrección de fallos y ejecución de pruebas.", "45 pruebas backend, 11 pruebas Flutter y análisis estático sin observaciones.", "Avanzado"),
]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
headers = ["Objetivo", "Actividad realizada", "Avance verificable", "Estado"]
widths = [0.55, 2.15, 3.35, 0.85]
for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
    cell.width = Inches(widths[i])
    cell.text = text
    set_cell_shading(cell, "D9D9D9")
    set_cell_margins(cell, 70, 70, 70, 70)
    format_paragraph(cell.paragraphs[0], size=8.5, bold=True, after=0)
set_repeat_table_header(table.rows[0])
for data in matrix_data:
    cells = table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, data)):
        cell.width = Inches(widths[i])
        cell.text = text
        set_cell_margins(cell, 70, 70, 70, 70)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        format_paragraph(cell.paragraphs[0], size=8.5, after=0)

add_heading(doc, "Criterio del 85 %", level=2)
add_body(doc, "El porcentaje se sustenta en la implementación de los flujos principales y en la corrección de los hallazgos de mayor impacto. No se declara el 100 % porque aún falta validación institucional en condiciones reales, preparación de despliegue y resolución de mejoras de producto pendientes.")

doc.add_page_break()
add_heading(doc, "3 Auditoría y pruebas ejecutadas")
add_body(doc, "La auditoría revisó el backend Laravel y la aplicación Flutter. Se corrigieron fallos de inventario, manejo de fechas, duplicación de lotes, concurrencia, bajas, seguridad, autorización, integridad de datos, sincronización y mensajes de validación.")
checks = [
    ("Pruebas backend", "45 de 45 aprobadas - 242 aserciones"),
    ("Pruebas Flutter", "11 de 11 aprobadas"),
    ("Análisis estático Flutter", "Sin observaciones"),
    ("Hallazgos críticos auditados", "Corregidos los cinco problemas de mayor impacto"),
    ("Seguridad e integridad", "Controles de permisos, límites, unicidad e índices reforzados"),
]
ct = doc.add_table(rows=0, cols=2)
ct.style = "Table Grid"
ct.alignment = WD_TABLE_ALIGNMENT.CENTER
for label, value in checks:
    row = ct.add_row()
    set_cell_shading(row.cells[0], "EDEDED")
    row.cells[0].text = label
    row.cells[1].text = value
    for idx, cell in enumerate(row.cells):
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        format_paragraph(cell.paragraphs[0], size=9, bold=(idx == 0), after=0)
add_heading(doc, "Pendientes reales", level=2)
add_body(doc, "Validación de campo con usuarios reales; revisión previa a la migración de producción; protección privada de fotografías; ajuste del modelo de alertas para reportes comparables; definición del alcance offline definitivo; y optimización de consultas de reportes.")

evidence = [
    ("01_inicio_sesion.png", "Inicio de sesión", "La autenticación restringe el acceso mediante número de ítem o contrato y contraseña."),
    ("02_pantalla_principal.png", "Pantalla principal", "Menú principal por rol con acceso a las funciones implementadas."),
    ("03_gestion_usuarios.png", "Gestión de usuarios", "Consulta de usuarios, roles, números de ítem y estados."),
    ("04_formulario_nuevo_usuario.png", "Formulario de usuarios", "Registro de cuentas con datos de identificación y asignación de rol."),
    ("05_catalogo_prendas.png", "Catálogo de prendas", "Administración de tipos de prendas utilizados por los procesos del sistema."),
    ("06_catalogo_areas.png", "Áreas y alias", "Áreas hospitalarias y alias empleados para la identificación y el OCR."),
    ("07_catalogo_plantillas.png", "Catálogo de plantillas", "Plantillas configuradas para los procesos de Salas y Quirófano."),
    ("08_impresion_plantillas.png", "Generación de plantillas", "Selección y generación de formularios imprimibles en PDF."),
    ("09_dashboard.png", "Dashboard de seguimiento", "Indicadores de alertas, lavado, bajas y prendas en circulación."),
]
for index, (filename, title_text, description) in enumerate(evidence, 1):
    doc.add_page_break()
    add_heading(doc, f"4.{index} {title_text}")
    add_body(doc, description)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(SHOTS / filename), width=Inches(7.0))
    cap = doc.add_paragraph(f"Captura {index}. {title_text} del sistema SICOTRAZ con datos de prueba.")
    format_paragraph(cap, size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)

doc.add_page_break()
add_heading(doc, "5 Conclusión del avance")
add_body(doc, "Las actividades realizadas mantienen coherencia con los objetivos del proyecto. El sistema ya demuestra los flujos centrales de digitalización, control y trazabilidad, y la auditoría incrementó su confiabilidad mediante correcciones verificadas por pruebas automatizadas.")
add_body(doc, "Por lo anterior, se registra un avance global del 85 %. El trabajo pendiente está delimitado y no se presenta como terminado: corresponde principalmente a validación operativa institucional, despliegue y mejoras posteriores a la prueba de campo.")

doc.core_properties.title = "Evidencia de avance 85 por ciento SICOTRAZ"
doc.core_properties.subject = "Objetivos actividades avances y evidencia"
doc.core_properties.author = "Pedro Jesus Chucamani Choque"
doc.save(OUTPUT)
print(OUTPUT)
